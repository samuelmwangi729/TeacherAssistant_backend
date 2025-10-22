from openpyxl import load_workbook
from io import BytesIO

def extract_xlsx(file_like):
    """Extracts formulas from a .xlsx file and returns a cleaned dictionary format."""
    data = file_like.read()
    wb = load_workbook(filename=BytesIO(data), data_only=False)  # keep formulas

    workbook = {'sheets': []}

    for ws in wb.worksheets:
        sheet_info = {'title': ws.title, 'formulas': []}

        for row in ws.iter_rows(values_only=False):
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith('='):
                    sheet_info['formulas'].append({
                        'address': cell.coordinate,
                        'formula': cell.value,
                        'value': cell.value  # original formula string
                    })

        workbook['sheets'].append(sheet_info)

    return workbook
